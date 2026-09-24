import datetime as dt
import io
from dataclasses import dataclass

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from app.backtest.config import BacktestConfig
from app.backtest.engine import Backtester
from app.backtest.metrics import compute_metrics
from app.exports import Table, backtest_tables, instances_table, live_trades_table, to_csv, to_xlsx
from app.live.stats import summarize
from app.strategies import registry
from app.ui.backtest_format import EXIT_REASON_LABELS, METRIC_GROUPS, METRIC_LABELS

UTC = dt.timezone.utc


def sample() -> Table:
    return Table("Prueba", ["Fecha (UTC)", "Texto", "Número", "Vacío", "Booleano"], [
        [dt.datetime(2026, 9, 24, 12, 30, 45, tzinfo=dt.timezone(dt.timedelta(hours=-3))), "áéí ñ, coma", 1.23456789012, None, True],
        [pd.Timestamp("2026-09-25 01:00", tz="UTC"), "=SUM(A1:A9)", float("nan"), None, False],
        [dt.datetime(2026, 9, 26), "-5 cmd", np.float64(2.5), None, None],
    ])


# ------------------------------------------------------------------ CSV

def test_csv_is_utf8_with_bom_and_keeps_accents_and_commas_quoted() -> None:
    raw = to_csv(sample())
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM: Excel muestra bien los acentos
    text = raw.decode("utf-8-sig")
    lines = text.split("\r\n")
    assert lines[0] == "Fecha (UTC),Texto,Número,Vacío,Booleano"
    assert lines[1].startswith('2026-09-24 12:30:45,"áéí ñ, coma",1.23456789012,,Sí')  # ya era hora argentina, coma escapada


def test_csv_neutralizes_formulas_and_drops_nan() -> None:
    lines = to_csv(sample()).decode("utf-8-sig").split("\r\n")
    assert "'=SUM(A1:A9)" in lines[2] and ",," in lines[2]      # formula anulada, NaN vacio
    assert "'-5 cmd" in lines[3] and "2.5" in lines[3]           # np.float64 se escribe como numero normal


# ------------------------------------------------------------------ Excel

def test_xlsx_has_one_sheet_per_table_typed_cells_header_and_filters() -> None:
    a, b = sample(), Table("Otra: hoja/con*chars?", ["x"], [[1], [2]])
    wb = load_workbook(io.BytesIO(to_xlsx([a, b, Table("Prueba", ["y"], [])])))
    assert wb.sheetnames == ["Prueba", "Otra  hoja con chars", "Prueba (2)"]  # sin caracteres invalidos y sin nombres repetidos
    ws = wb["Prueba"]
    assert [c.value for c in ws[1]] == a.headers and ws["A1"].font.bold and ws.freeze_panes == "A2"
    assert ws["A2"].value == dt.datetime(2026, 9, 24, 12, 30, 45)   # fecha real, en hora argentina y sin zona
    assert ws["A3"].value == dt.datetime(2026, 9, 24, 22, 0)         # 01:00 UTC del 25/09 = 22:00 del 24/09 en Argentina
    assert ws["C2"].value == pytest.approx(1.23456789012)            # numero real, no texto ni redondeado
    assert ws["B3"].value == "'=SUM(A1:A9)" and ws["C3"].value is None
    assert ws["E2"].value == "Sí"
    assert ws.auto_filter.ref == "A1:E4"


def test_sheet_name_is_limited_to_31_characters() -> None:
    wb = load_workbook(io.BytesIO(to_xlsx([Table("x" * 80, ["a"], [[1]])])))
    assert len(wb.sheetnames[0]) == 31


# ------------------------------------------------------------------ operativa en vivo

@dataclass
class T:
    id: int
    strategy_instance_id: int
    symbol: str
    side: str
    qty: float
    entry_price: float
    exit_price: float | None
    stop_loss: float | None
    take_profit: float | None
    pnl: float | None
    exit_reason: str | None
    opened_at: dt.datetime
    closed_at: dt.datetime | None


def test_live_trades_table_computes_duration_and_state() -> None:
    opened = dt.datetime(2026, 9, 24, 0, 0, tzinfo=UTC)
    trades = [T(1, 7, "BTCUSDT", "long", 0.003, 84000.0, 85000.0, 82000.0, None, 2.5, "senal", opened, opened + dt.timedelta(hours=3)),
              T(2, 7, "BTCUSDT", "short", 0.003, 84000.0, None, None, None, None, None, opened, None)]
    table = live_trades_table(trades, {7: "rsi-1"})
    assert table.rows[0][1] == "rsi-1" and table.rows[0][13] == pytest.approx(3.0) and table.rows[0][14] == "Cerrado"
    assert table.rows[1][13] is None and table.rows[1][14] == "Abierto"
    assert len(table.headers) == len(table.rows[0])


def test_instances_table_matches_the_stats() -> None:
    @dataclass
    class I:
        id: int = 1
        name: str = "a"
        strategy_key: str = "rsi_reversion"
        symbol: str = "BTCUSDT"
        timeframe: str = "15"
        initial_capital: float = 1000.0
        is_active: bool = False
        params: dict = None

    opened = dt.datetime(2026, 9, 24, tzinfo=UTC)
    trades = {1: [T(1, 1, "BTCUSDT", "long", 1, 1, 2, None, None, 6.0, "x", opened, opened), T(2, 1, "BTCUSDT", "long", 1, 1, 2, None, None, -3.0, "x", opened, opened)]}
    table = instances_table([I(params={"a": 1})], trades, summarize)
    row = dict(zip(table.headers, table.rows[0], strict=True))
    assert row["PnL realizado (USD)"] == 3.0 and row["Win rate (%)"] == 50.0 and row["Profit factor"] == 2.0
    assert row["Estado"] == "Apagada"


# ------------------------------------------------------------------ backtest

def _backtest():
    rng = np.random.default_rng(2)
    n = 1500
    closes = 100 * np.exp(np.cumsum(rng.normal(0, 0.004, n)))
    idx = pd.date_range("2026-01-01", periods=n, freq="1h", tz="UTC")
    opens = np.concatenate([[closes[0]], closes[:-1]])
    candles = pd.DataFrame({"open": opens, "high": np.maximum(opens, closes) * 1.002, "low": np.minimum(opens, closes) * 0.998, "close": closes, "volume": 1.0}, index=idx)
    cls = registry.get("donchian_breakout")
    cfg = BacktestConfig(timeframe="60")
    result = Backtester(execution=cfg.execution, risk=cfg.risk).run(cls(cls.params_model()), candles, 1000.0)
    result.symbol, result.timeframe, result.config = "TESTUSDT", "60", cfg.to_dict()
    return result, compute_metrics(result, "60"), cfg


def test_backtest_export_reconciles_with_the_result_and_survives_excel() -> None:
    result, metrics, cfg = _backtest()
    assert len(result.trades) > 5
    tables = backtest_tables(result, metrics, cfg.to_dict(), "Donchian", "v1+abc", METRIC_GROUPS, METRIC_LABELS, EXIT_REASON_LABELS)
    by_name = {t.name: t for t in tables}
    assert {"Métricas", "Configuración", "Operaciones", "Equity"} <= set(by_name)

    trades = by_name["Operaciones"]
    idx = trades.headers.index("PnL neto")
    assert sum(r[idx] for r in trades.rows) == pytest.approx(result.equity_curve.iloc[-1] - result.initial_capital)  # las cuentas cierran
    assert all(len(r) == len(trades.headers) for r in trades.rows)

    equity = by_name["Equity"]
    assert len(equity.rows) == len(result.equity_curve) and min(r[2] for r in equity.rows) <= 0 and max(r[2] for r in equity.rows) <= 0

    m = {r[1]: r[2] for r in by_name["Métricas"].rows}
    assert m[METRIC_LABELS["total_return_pct"]] == pytest.approx(metrics["total_return_pct"])

    wb = load_workbook(io.BytesIO(to_xlsx(tables)))          # el archivo completo se abre sin errores
    assert wb["Operaciones"].max_row == len(result.trades) + 1
    assert wb["Equity"].max_row == len(result.equity_curve) + 1
    assert to_csv(trades).count(b"\r\n") == len(result.trades) + 1


def test_instants_are_exported_in_argentina_time_with_that_in_the_header() -> None:
    utc = dt.datetime(2026, 9, 24, 3, 0, tzinfo=dt.timezone.utc)           # como se guarda: UTC
    table = live_trades_table([T(1, 7, "BTCUSDT", "long", 1, 100.0, 110.0, None, None, 5.0, "senal", utc, utc + dt.timedelta(hours=2))], {7: "a"})
    row = dict(zip(table.headers, table.rows[0], strict=True))
    from app.exports import _cell  # la conversion ocurre al escribir la celda (la tabla conserva el instante original)

    assert _cell(row["Apertura (hora argentina)"]) == dt.datetime(2026, 9, 24, 0, 0)   # 03:00 UTC = 00:00 en Argentina
    assert _cell(row["Cierre (hora argentina)"]) == dt.datetime(2026, 9, 24, 2, 0)
    assert not any("UTC" in h for h in table.headers)
    assert to_csv(table).decode("utf-8-sig").splitlines()[1].count("2026-09-24 00:00:00") == 1
