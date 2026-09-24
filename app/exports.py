"""Exportacion de datos a Excel (.xlsx) y CSV. Es puro (sin interfaz ni base): arma tablas y las serializa.

Convenciones:
- Los instantes se exportan en la hora local (Argentina), sin zona horaria (Excel no admite fechas con zona) y con
  «(hora argentina)» en el titulo. Internamente todo se guarda en UTC.
- Los numeros salen con toda su precision, no redondeados como en pantalla: para analizar despues.
- Los textos que empiezan con = + - @ se prefijan con una comilla: evita que Excel los ejecute como formulas.
- El CSV va en UTF-8 con BOM (Excel muestra bien los acentos), separado por comas y con punto decimal."""

import csv
import datetime as dt
import io
import json
import re
from dataclasses import dataclass, field

import numpy as np

from app.timeutil import naive_local


@dataclass
class Table:
    name: str
    headers: list[str]
    rows: list[list] = field(default_factory=list)


# ------------------------------------------------------------------ serializacion

def _cell(value):
    """Valor listo para escribir en una celda (igual para xlsx y csv)."""
    if value is None:
        return None
    if isinstance(value, (np.floating, np.integer)):
        value = value.item()
    if isinstance(value, float) and (value != value or value in (float("inf"), float("-inf"))):
        return None  # NaN / inf no existen en una planilla
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, dt.datetime):
        return naive_local(value).replace(microsecond=0)
    if hasattr(value, "to_pydatetime"):  # pd.Timestamp
        return _cell(value.to_pydatetime())
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, default=str)
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(table: Table) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerow(table.headers)
    for row in table.rows:
        writer.writerow(["" if (c := _cell(v)) is None else (c.strftime("%Y-%m-%d %H:%M:%S") if isinstance(c, dt.datetime) else c) for v in row])
    return ("﻿" + buffer.getvalue()).encode("utf-8")


_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def _sheet_names(tables: list[Table]) -> list[str]:
    names, used = [], set()
    for table in tables:
        base = (_INVALID_SHEET_CHARS.sub(" ", table.name).strip() or "Hoja")[:31]
        name, n = base, 2
        while name.lower() in used:
            suffix = f" ({n})"
            name, n = base[: 31 - len(suffix)] + suffix, n + 1
        used.add(name.lower())
        names.append(name)
    return names


def to_xlsx(tables: list[Table]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font, header_fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="2A78D6")
    for table, name in zip(tables, _sheet_names(tables), strict=True):
        sheet = workbook.create_sheet(name)
        sheet.append(table.headers)
        for cell in sheet[1]:
            cell.font, cell.fill, cell.alignment = header_font, header_fill, Alignment(vertical="center", wrap_text=True)
        widths = [len(str(h)) for h in table.headers]
        for row in table.rows:
            values = [_cell(v) for v in row]
            sheet.append(values)
            for i, v in enumerate(values):
                widths[i] = max(widths[i], len(v.strftime("%Y-%m-%d %H:%M") if isinstance(v, dt.datetime) else str(v if v is not None else "")))
        for i, cells in enumerate(sheet.iter_cols(min_row=2, max_row=sheet.max_row), start=1):
            for cell in cells:
                if isinstance(cell.value, dt.datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        sheet.freeze_panes = "A2"
        if table.rows:
            sheet.auto_filter.ref = sheet.dimensions
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------ operativa en vivo

def live_trades_table(trades, names: dict[int, str]) -> Table:
    rows = []
    for t in trades:
        duration = (t.closed_at - t.opened_at).total_seconds() / 3600 if t.closed_at else None
        rows.append([
            t.id, names.get(t.strategy_instance_id, ""), t.symbol, t.side, t.qty, t.entry_price, t.exit_price,
            t.stop_loss, t.take_profit, t.pnl, t.exit_reason, t.opened_at, t.closed_at, duration,
            "Cerrado" if t.closed_at else "Abierto",
        ])
    return Table("Trades", [
        "ID", "Instancia", "Símbolo", "Lado", "Cantidad", "Precio entrada", "Precio salida", "Stop loss", "Take profit",
        "PnL (USD)", "Motivo de salida", "Apertura (hora argentina)", "Cierre (hora argentina)", "Duración (horas)", "Estado",
    ], rows)


def live_orders_table(orders, names: dict[int, str]) -> Table:
    return Table("Órdenes", [
        "ID", "Fecha (hora argentina)", "Instancia", "Símbolo", "Lado", "Tipo", "Cantidad", "Precio", "Estado", "ID en Bybit",
    ], [[o.id, o.created_at, names.get(o.strategy_instance_id, ""), o.symbol, o.side, o.order_type, o.qty, o.price,
         o.status, o.exchange_order_id] for o in orders])


def instances_table(instances, trades_by_instance: dict[int, list], summarize) -> Table:
    rows = []
    for i in instances:
        s = summarize(trades_by_instance.get(i.id, []))
        rows.append([
            i.name, i.strategy_key, i.symbol, i.timeframe, i.initial_capital, "Activa" if i.is_active else "Apagada",
            s["closed"], s["open"], s["total_pnl"], s["win_rate_pct"], s["profit_factor"], i.params,
        ])
    return Table("Instancias", [
        "Instancia", "Estrategia", "Símbolo", "Timeframe", "Capital (USD)", "Estado", "Trades cerrados", "Trades abiertos",
        "PnL realizado (USD)", "Win rate (%)", "Profit factor", "Parámetros",
    ], rows)


# ------------------------------------------------------------------ backtests

def _flatten(data: dict, prefix: str = "") -> list[tuple[str, object]]:
    items = []
    for key, value in (data or {}).items():
        name = f"{prefix}{key}"
        if isinstance(value, dict):
            items.extend(_flatten(value, name + "."))
        else:
            items.append((name, value))
    return items


def backtest_tables(result, metrics: dict, config: dict, strategy_name: str, strategy_version: str,
                    metric_groups: dict, metric_labels: dict, exit_labels: dict) -> list[Table]:
    """Hojas: Métricas, Avisos (si hay), Configuración, Operaciones y Equity."""
    tables = []
    rows = [[group, metric_labels.get(key, key), metrics.get(key)] for group, keys in metric_groups.items() for key in keys if key in metrics]
    tables.append(Table("Métricas", ["Grupo", "Métrica", "Valor"], rows))
    warnings = metrics.get("warnings") or []
    if warnings:
        tables.append(Table("Avisos", ["Nivel", "Aviso"], [[w["level"], w["text"]] for w in warnings]))

    info = [["Estrategia", strategy_name], ["Versión de la estrategia", strategy_version], ["Símbolo", result.symbol],
            ["Timeframe", result.timeframe], ["Capital inicial", result.initial_capital]]
    info += [[f"Parámetros · {k}", v] for k, v in (result.params or {}).items()]
    info += [[k, v] for k, v in _flatten(config)]
    tables.append(Table("Configuración", ["Parámetro", "Valor"], info))

    tables.append(Table("Operaciones", [
        "ID", "Entrada (hora argentina)", "Salida (hora argentina)", "Lado", "Precio entrada", "Precio salida", "Cantidad", "Nocional (USD)",
        "PnL bruto", "Comisión entrada", "Comisión salida", "Funding", "Costo slippage", "PnL neto", "Retorno neto (%)",
        "Capital antes", "Duración (horas)", "Stop loss", "Take profit", "Salida por", "Motivo de entrada",
        "Cerrada al final del backtest", "Limitada por apalancamiento",
    ], [[
        t.id, t.entry_time, t.exit_time, t.side, t.entry_price, t.exit_price, t.qty, t.notional, t.gross_pnl, t.entry_fee,
        t.exit_fee, t.funding, t.slippage_cost, t.pnl, None if t.pnl_pct is None else t.pnl_pct * 100, t.equity_before,
        None if t.duration is None else t.duration.total_seconds() / 3600, t.stop_loss, t.take_profit,
        exit_labels.get(t.exit_reason, t.exit_reason), t.reason, t.open_at_end, t.capped_by_leverage,
    ] for t in result.trades]))

    equity = result.equity_curve
    values = equity.to_numpy(dtype=float)
    peak = np.maximum.accumulate(np.concatenate(([result.initial_capital], values)))[1:]
    drawdown = (values / peak - 1.0) * 100
    tables.append(Table("Equity", ["Fecha (hora argentina)", "Capital (USD)", "Drawdown (%)"],
                        [[ts, v, d] for ts, v, d in zip(equity.index, values, drawdown, strict=True)]))
    return tables


def runs_table(runs, timeframe_labels: dict) -> Table:
    rows = []
    for r in runs:
        m = r.metrics or {}
        rows.append([r.created_at, r.strategy_key, r.symbol, timeframe_labels.get(r.timeframe, r.timeframe),
                     r.start_date, r.end_date, m.get("total_return_pct"), m.get("sharpe_ratio"), m.get("max_drawdown_pct"),
                     m.get("num_trades"), r.params])
    return Table("Backtests", ["Fecha (hora argentina)", "Estrategia", "Símbolo", "Timeframe", "Desde (hora argentina)", "Hasta (hora argentina)",
                               "Retorno (%)", "Sharpe", "Drawdown máx. (%)", "Operaciones", "Parámetros"], rows)


# ------------------------------------------------------------------ costos y stress tests

_STRESS_HEADERS = ["Retorno (%)", "Sharpe", "Drawdown máx. (%)", "Profit factor", "Capital final (USD)", "Operaciones", "Expectativa (USD)"]
_STRESS_KEYS = ["total_return_pct", "sharpe_ratio", "max_drawdown_pct", "profit_factor", "final_equity", "num_trades", "expectancy"]


def stress_table(results) -> Table:
    base = next((r.metrics for r in results if r.key == "base" and r.metrics), None)
    rows = []
    for r in results:
        m = r.metrics or {}
        delta = None if not base or not m else m.get("total_return_pct", 0) - base.get("total_return_pct", 0)
        rows.append([r.label, r.description, *[m.get(k) for k in _STRESS_KEYS], delta, r.error])
    return Table("Stress tests", ["Escenario", "Descripción", *_STRESS_HEADERS, "Δ retorno vs base (pp)", "Error"], rows)


def cost_tables(cost) -> list[Table]:
    def points(name: str, x_header: str, pts) -> Table:
        return Table(name, [x_header, *_STRESS_HEADERS, "Error"],
                     [[p.label if p.x is None else p.x, *[(p.metrics or {}).get(k) for k in _STRESS_KEYS], p.error] for p in pts])

    heat = Table("Slippage x comisiones", ["Comisiones (× las actuales) / Slippage (bps)", *cost.heat_x],
                 [[y, *row] for y, row in zip(cost.heat_y, cost.heat_z, strict=True)])
    summary = Table("Punto de equilibrio", ["Costo", "Valor de equilibrio", "Nota"], [
        ["Slippage (bps)", cost.break_even_slippage[0], cost.break_even_slippage[1]],
        ["Comisiones (× las actuales)", cost.break_even_fees[0], cost.break_even_fees[1]],
    ])
    return [summary, points("Slippage", "Slippage (bps)", cost.slippage), points("Comisiones", "Comisiones (× las actuales)", cost.fees),
            points("Funding", "Escenario de funding", cost.funding), heat]
